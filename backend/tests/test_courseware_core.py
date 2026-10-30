"""Synthetic fixtures test mechanics only; they are never registered as lessons."""
from copy import deepcopy
import hashlib
import os
from pathlib import Path
import subprocess

import pytest

from review.blocks import speech_blocks
from review.courseware_alignment import align_document
from review.courseware_documents import CoursewareError, import_document, _native_text_quality
from review.courseware_terms import revise_transcript


def _pdf(path, texts):
    import pymupdf
    with pymupdf.open() as document:
        for text in texts:
            page = document.new_page(width=720, height=405)
            page.insert_text((45, 70), text, fontsize=22)
        document.save(path)


def _document(texts):
    return {'document_id': 'Dsynthetic', 'pages': [
        {'id': f'Psynthetic_{index}', 'document_id': 'Dsynthetic', 'page_number': index, 'text': text}
        for index, text in enumerate(texts, 1)]}


def _bundle(text='這裏我們使用鄭賢定理，然後討論相量。'):
    transcript = [{'id': 'T0001', 'text': text, 'start_sec': 0, 'end_sec': 4},
                  {'id': 'T0002', 'text': '此處保留完整上下文。', 'start_sec': 4, 'end_sec': 8},
                  {'id': 'T0003', 'text': '下一個未修改的證據區塊。', 'start_sec': 25, 'end_sec': 30}]
    return {'transcript': transcript, 'slides': [{'id': 'S0001', 'start_sec': 0, 'end_sec': 10}],
            'evidence_blocks': speech_blocks(transcript)}


def _matched(document, status='auto_matched'):
    document['matches'] = [{'id': 'Msynthetic_S0001', 'document_id': document['document_id'],
                            'slide_id': 'S0001', 'page_id': document['pages'][0]['id'],
                            'start_sec': 0, 'end_sec': 10, 'status': status,
                            'score': 0.95, 'method': 'synthetic_fixture_match'}]
    return document


def test_actual_pdf_parse_render_source_preservation_and_cache(tmp_path):
    from PIL import Image
    source = tmp_path / 'synthetic-source.pdf'
    _pdf(source, ['SYNTHETIC TEST: vector arithmetic', 'SYNTHETIC TEST: sine theorem'])
    original = source.read_bytes()
    result = import_document(source, tmp_path / 'lesson', 'test-courseware')
    assert len(result['pages']) == 2
    assert result['sha256'] == hashlib.sha256(original).hexdigest()
    assert (tmp_path / 'lesson' / result['source_path']).read_bytes() == original
    assert source.read_bytes() == original
    assert 'vector arithmetic' in result['pages'][0]['text']
    assert result['pages'][0]['text_source'] == 'pdf_native'
    with Image.open(tmp_path / 'lesson' / result['pages'][0]['image_path']) as image:
        assert image.width > 500 and image.width * image.height <= 3_000_000
        assert len(image.convert('L').getcolors(256) or []) > 1
    assert result['pages'][0]['image_url'].startswith('/api/review/media/test-courseware/courseware/')
    assert import_document(source, tmp_path / 'lesson', 'test-courseware') == result


def test_scan_uses_real_ocr_fallback(tmp_path):
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont
    source = tmp_path / 'synthetic-scan.pdf'
    picture = Image.new('RGB', (1500, 600), 'white')
    draw = ImageDraw.Draw(picture)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    if not font_path.exists():
        pytest.skip('Windows test font is unavailable')
    font = ImageFont.truetype(str(font_path), 70)
    draw.text((80, 150), 'SYNTHETIC VECTOR ARITHMETIC', font=font, fill='black')
    png = tmp_path / 'scan.png'
    picture.save(png)
    with pymupdf.open() as document:
        page = document.new_page(width=750, height=300)
        page.insert_image(page.rect, filename=str(png))
        document.save(source)
    result = import_document(source, tmp_path / 'lesson', 'test-scan')
    assert result['pages'][0]['text_source'] == 'tesseract_chi_sim_eng'
    assert 'VECTOR' in result['pages'][0]['text'].upper()
    assert not result['pages'][0]['native_text'].strip()


def test_native_quality_rejects_corrupt_font_map():
    assert _native_text_quality('Good readable native PDF text') == 'usable'
    assert _native_text_quality('1URDU>PPFQD\x06>VVFDQPBQ]V\x06>\x06?F]') == 'broken_font_map'
    assert _native_text_quality('') == 'too_short'


def test_invalid_source_and_page_limit_are_rejected_without_artifacts(tmp_path):
    source = tmp_path / 'not-a-pdf.pdf'
    source.write_bytes(b'<html>download failed</html>')
    with pytest.raises(CoursewareError, match='不是 PDF'):
        import_document(source, tmp_path / 'lesson', 'invalid-source')
    _pdf(source, ['x'] * 201)
    with pytest.raises(CoursewareError, match='200'):
        import_document(source, tmp_path / 'lesson', 'too-many-pages')
    assert not list((tmp_path / 'lesson' / 'courseware').iterdir())


def test_repeated_and_skipped_pages_keep_observed_timing(tmp_path):
    document = _document(['正弦定理适用于三角形各边与对应角关系', '平面向量加法满足平行四边形法则', '空间坐标系与解析几何基础'])
    slides = [{'id': 'S1', 'text': document['pages'][1]['text'], 'start_sec': 2, 'end_sec': 6},
              {'id': 'S2', 'text': document['pages'][0]['text'], 'start_sec': 9, 'end_sec': 12},
              {'id': 'S3', 'text': document['pages'][1]['text'], 'start_sec': 12, 'end_sec': 20}]
    result = align_document(document, {'slides': slides}, tmp_path)
    assert [row['page_id'] for row in result] == ['Psynthetic_2', 'Psynthetic_1', 'Psynthetic_2']
    assert [(row['start_sec'], row['end_sec']) for row in result] == [(2, 6), (9, 12), (12, 20)]
    assert all(row['status'] == 'auto_matched' for row in result)


def test_duplicate_pages_abstain_and_empty_observation_never_gets_time(tmp_path):
    text = '正弦定理适用于三角形各边与对应角关系'
    document = _document([text, text])
    slides = [{'id': 'S1', 'text': text, 'start_sec': 3, 'end_sec': 9},
              {'id': 'S2', 'text': '', 'start_sec': 20, 'end_sec': 35}]
    matches = align_document(document, {'slides': slides}, tmp_path)
    assert matches[0]['status'] == 'ambiguous' and matches[0]['margin'] == 0
    assert matches[1]['status'] == 'unmatched'
    assert all(row['page_id'] is None for row in matches)


def test_acceptance_threshold_does_not_erase_close_runner_up(tmp_path, monkeypatch):
    import review.courseware_alignment as alignment
    document = _document(['candidate one', 'candidate two'])
    values = {'candidate one': 0.430529, 'candidate two': 0.429333}
    monkeypatch.setattr(alignment, 'lexical_score', lambda observed, page: values[page])
    matches = alignment._align_in_process(document, {'slides': [
        {'id': 'S1', 'text': 'observed text', 'start_sec': 0, 'end_sec': 4}]}, tmp_path)
    assert matches[0]['status'] == 'ambiguous'
    assert matches[0]['page_id'] is None and matches[0]['margin'] < 0.002


def test_actual_render_keypoints_and_blank_frame_abstain(tmp_path):
    import cv2
    import numpy as np
    # Deterministic geometric pattern tests the matcher, not source authenticity.
    image = np.full((650, 1000, 3), 255, dtype=np.uint8)
    rng = np.random.default_rng(55)
    for index in range(70):
        point = tuple(int(value) for value in rng.integers([30, 30], [950, 600]))
        cv2.putText(image, str(index), point, cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    def save(name, picture):
        cv2.imencode('.png', picture)[1].tofile(tmp_path / name)
    save('page.png', image)
    save('observation.png', cv2.resize(image, (700, 455)))
    save('blank.png', np.full((500, 700, 3), 255, dtype=np.uint8))
    document = _document([''])
    document['pages'][0]['image_path'] = 'page.png'
    matches = align_document(document, {'slides': [
        {'id': 'S1', 'text': '', 'image_path': 'observation.png', 'start_sec': 0, 'end_sec': 3},
        {'id': 'S2', 'text': '', 'image_path': 'blank.png', 'start_sec': 3, 'end_sec': 6}]}, tmp_path)
    assert matches[0]['status'] == 'auto_matched'
    assert matches[0]['method'] == 'observed_frame_orb_ransac'
    assert matches[1]['status'] == 'unmatched'


@pytest.mark.parametrize('status', ['auto_matched', 'confirmed'])
def test_source_grounded_homophones_preserve_raw_script_context_and_ids(status):
    bundle = _bundle()
    original = deepcopy(bundle)
    document = _matched(_document(['正弦定理 平面向量']), status)
    result = revise_transcript(bundle, document)
    assert bundle == original
    revision = result['transcript_revisions'][0]
    assert revision['corrected_text'] == '這裏我們使用正弦定理，然後討論向量。'
    assert [edit['original'] for edit in revision['edits']] == ['鄭賢定理', '相量']
    assert all(edit['page_ids'] and edit['match_ids'] for edit in revision['edits'])
    blocks = result['revised_evidence_blocks']
    assert blocks[0]['source_refs'] == ['T0001', 'T0002']
    assert blocks[0]['end_sec'] == 8 and '完整上下文' in blocks[0]['text']
    assert blocks[1] == bundle['evidence_blocks'][1]
    assert result['courseware_evidence'][0]['match_id'] == document['matches'][0]['id']
    assert result['courseware_evidence'][0]['status'] == status
    assert result == revise_transcript(bundle, document)


def test_unmatched_missing_page_or_forged_timing_cannot_correct():
    bundle = _bundle()
    document = _matched(_document(['正弦定理 平面向量']), 'ambiguous')
    assert not revise_transcript(bundle, document)['transcript_revisions']
    document['matches'][0].update(status='auto_matched', start_sec=1)
    assert not revise_transcript(bundle, document)['courseware_evidence']
    document['matches'][0].update(start_sec=0, page_id='P_nonexistent')
    assert not revise_transcript(bundle, document)['transcript_revisions']


def test_numbers_negation_formulas_and_traditional_script_remain():
    bundle = _bundle('不是餘弦定理，2+3=5；相量=x，三個相量，鄭賢定理。')
    result = revise_transcript(bundle, _matched(_document(['余弦定理 正弦定理 向量'])))
    corrected = result['transcript_revisions'][0]['corrected_text']
    assert '不是餘弦定理，2+3=5；相量=x' in corrected
    assert '三個向量' in corrected and '正弦定理' in corrected


@pytest.mark.parametrize('formula', ['相量 = x', '相量²', '相量−x', 'x = 相量', '相量 + 2',
                                     '相量-x', '(相量)=x', '相量＝x', '（相量）－ｘ', 'x + (相量)'])
def test_formula_adjacent_homophones_abstain(formula):
    result = revise_transcript(_bundle(formula), _matched(_document(['向量'])))
    assert not result['transcript_revisions']


def test_formula_guard_preserves_natural_language_corrections_in_separate_clause():
    text = '相量＝x；根据相量我们证明正弦定理。'
    result = revise_transcript(_bundle(text), _matched(_document(['向量 正弦定理'])))
    assert result['transcript_revisions'][0]['corrected_text'] == '相量＝x；根据向量我们证明正弦定理。'


def test_conflicting_aligned_canonical_terms_abstain():
    bundle = _bundle('現在討論相量。')
    document = _matched(_document(['向量和相量的区别']))
    assert not revise_transcript(bundle, document)['transcript_revisions']


@pytest.mark.skipif(os.environ.get('COURSEWARE_TEST_POWERPOINT') != '1',
                    reason='Explicit local PowerPoint smoke; opens only a synthetic temporary deck')
@pytest.mark.parametrize('source_format', ['pptx', 'ppt'])
def test_powerpoint_real_export_smoke(tmp_path, source_format):
    from pptx import Presentation
    from pptx.util import Inches
    source = tmp_path / 'SYNTHETIC-not-a-real-lesson.pptx'
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = 'SYNTHETIC TEST DECK: source-preserving PowerPoint export'
    presentation.save(source)
    if source_format == 'ppt':
        converter = Path(__file__).resolve().parents[1] / 'scripts' / 'convert_courseware.ps1'
        fixture_converter = tmp_path / 'make-synthetic-legacy-fixture.ps1'
        fixture_converter.write_text(converter.read_text(encoding='utf-8').replace(
            '$presentation.SaveAs($outputResolved, 32)', '$presentation.SaveAs($outputResolved, 1)'), encoding='utf-8')
        legacy = source.with_suffix('.ppt')
        converted = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                                    '-File', str(fixture_converter), '-InputPath', str(source),
                                    '-OutputPath', str(legacy), '-OwnerPath', str(tmp_path / 'fixture-owner.json')],
                                   capture_output=True, timeout=120, creationflags=subprocess.CREATE_NO_WINDOW)
        assert converted.returncode == 0, converted.stderr.decode('utf-8', errors='replace')
        assert legacy.read_bytes().startswith(bytes.fromhex('D0CF11E0A1B11AE1'))
        source = legacy
    original = source.read_bytes()
    result = import_document(source, tmp_path / 'lesson', 'synthetic-ppt-smoke')
    assert result['format'] == source_format
    assert 'PowerPoint export' in result['pages'][0]['text']
    assert result['provenance']['converter'] == 'microsoft_powerpoint_readonly_pdf_export'
    assert (tmp_path / 'lesson' / result['pages'][0]['image_path']).stat().st_size > 1000
    assert source.read_bytes() == original
