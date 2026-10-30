"""Read-only courseware ingestion. Rendered pages always come from the source PDF.

PowerPoint requires an unused desktop PowerPoint installation. We fail closed if
an existing PowerPoint process makes ownership of COM automation uncertain.
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unicodedata
from urllib.parse import quote
import zipfile

MAX_DOCUMENT_BYTES = 64 * 1024 * 1024
MAX_PAGES = 200
MAX_PAGE_PIXELS = 3_000_000
MAX_TOTAL_PIXELS = 220_000_000
RENDER_LONG_EDGE = 1800
CONVERT_TIMEOUT_SEC = 120
IMPORTER_VERSION = 'v3'


class CoursewareError(ValueError):
    """A source could not be safely imported; no substitute pages were created."""


def _hash(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_source(path):
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_DOCUMENT_BYTES:
        raise CoursewareError('The courseware file is empty, missing, or larger than 64 MiB.')
    extension = path.suffix.lower()
    if extension not in {'.pdf', '.pptx', '.ppt'}:
        raise CoursewareError('Only PDF, PPTX, and PPT courseware files are supported.')
    with path.open('rb') as handle:
        signature = handle.read(1024)
    if extension == '.pdf' and b'%PDF-' not in signature:
        raise CoursewareError('The file content is not PDF.')
    if extension == '.ppt' and not signature.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        raise CoursewareError('The file content is not legacy PPT.')
    if extension == '.pptx':
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 20000 or sum(x.file_size for x in members) > 512 * 1024 * 1024:
                    raise CoursewareError('The expanded PPTX size exceeds the processing limit.')
                names = set(archive.namelist())
                if 'ppt/presentation.xml' not in names or '[Content_Types].xml' not in names:
                    raise CoursewareError('The file content is not PPTX.')
                if any('vbaproject' in name.lower() for name in names):
                    raise CoursewareError('PPTX files containing macros are not supported. Save the file as PDF.')
        except zipfile.BadZipFile as exc:
            raise CoursewareError('The PPTX archive is corrupt.') from exc
    return extension[1:]


def _convert_powerpoint(source, destination, work_dir):
    if os.name != 'nt':
        raise CoursewareError('Rendering PPT/PPTX requires local Windows PowerPoint. You can also upload a PDF.')
    executable = shutil.which('powershell') or shutil.which('pwsh')
    script = Path(__file__).resolve().parents[1] / 'scripts' / 'convert_courseware.ps1'
    if not executable or not script.is_file():
        raise CoursewareError('The PowerPoint converter is unavailable. Upload a PDF.')
    owner_path = work_dir / 'converter-owner.json'
    args = [executable, '-NoLogo', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
            '-File', str(script), '-InputPath', str(source), '-OutputPath', str(destination),
            '-OwnerPath', str(owner_path), '-MaxPages', str(MAX_PAGES)]
    try:
        result = subprocess.run(args, capture_output=True, text=True, encoding='utf-8', errors='replace',
                                timeout=CONVERT_TIMEOUT_SEC, creationflags=subprocess.CREATE_NO_WINDOW)
    except subprocess.TimeoutExpired as exc:
        # The script records only a verified newly created automation process.
        try:
            subprocess.run([executable, '-NoProfile', '-NonInteractive', '-File', str(script),
                            '-OwnerPath', str(owner_path), '-CleanupOnly'], capture_output=True,
                           timeout=15, creationflags=subprocess.CREATE_NO_WINDOW)
        except (subprocess.SubprocessError, OSError):
            pass # An uninspectable desktop instance must remain untouched.
        raise CoursewareError('PowerPoint conversion timed out. No substitute courseware was generated. Save as PDF and try again.') from exc
    finally:
        # Internal process identifiers must not become public artifact metadata.
        owner_path.unlink(missing_ok=True)
    if result.returncode or not destination.is_file():
        detail = (result.stderr or result.stdout).strip()[-1800:]
        raise CoursewareError('PowerPoint conversion failed: ' + detail)
    return 'microsoft_powerpoint_readonly_pdf_export'


def _pptx_text(path):
    from pptx import Presentation
    deck = Presentation(str(path))
    if not 0 < len(deck.slides) <= MAX_PAGES:
        raise CoursewareError('Courseware must contain 1–200 pages.')
    def texts(shapes):
        rows = []
        for shape in shapes:
            if shape.has_text_frame:
                rows.append(shape.text_frame.text)
            if shape.has_table:
                rows.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if hasattr(shape, 'shapes'):
                rows.extend(texts(shape.shapes))
        return rows
    return ['\n'.join(texts(slide.shapes)).strip() for slide in deck.slides]


def _ocr(image_path):
    import pytesseract
    from PIL import Image
    with Image.open(image_path) as picture:
        return pytesseract.image_to_string(picture, lang='chi_sim+eng',
                                          config='--oem 1 --psm 6', timeout=30).strip()


def _native_text_quality(text):
    """Reject broken font maps as well as empty scanned-page text."""
    compact = ''.join(text.split())
    if len(compact) < 8:
        return 'too_short'
    controls = sum(unicodedata.category(char) in {'Cc', 'Co'} and char not in '\t\r\n' for char in text)
    replacement = text.count('\ufffd')
    if replacement or controls >= 2 or controls / max(len(text), 1) > 0.015:
        return 'broken_font_map'
    return 'usable'


def import_document(path, lesson_dir, lesson_id):
    """Preserve the uploaded bytes; create genuine PNG pages and native/OCR text.

    Does not change bundle.json or any original transcript. Paths in returned
    metadata are lesson-relative; original local source paths are not exposed.
    """
    from .common import valid_id
    valid_id(lesson_id)
    source, lesson_dir = Path(path).resolve(), Path(lesson_dir).resolve()
    kind = _validate_source(source)
    digest = _hash(source)
    document_id = 'D' + digest[:16]
    # Versioned immutable render folders keep earlier source artifacts intact
    # when native-text quality detection changes.
    relative_root = Path('courseware') / (digest[:16] + '_' + IMPORTER_VERSION)
    root = lesson_dir / relative_root
    parent = root.parent
    parent.mkdir(parents=True, exist_ok=True)
    metadata_path = root / 'document.json'
    if metadata_path.is_file():
        cached = json.loads(metadata_path.read_text(encoding='utf-8'))
        if cached.get('sha256') != digest:
            raise CoursewareError('Courseware storage ID conflict.')
        paths = [root / ('original.' + kind), *[lesson_dir / row['image_path'] for row in cached['pages']]]
        if all(item.is_file() for item in paths) and _hash(paths[0]) == digest:
            return cached
        raise CoursewareError('The existing courseware cache is incomplete. Rebuild the courseware artifact directory.')
    stage = Path(tempfile.mkdtemp(prefix='.' + digest[:16] + '-', dir=parent))
    try:
        original = stage / ('original.' + kind)
        shutil.copyfile(source, original)
        if _hash(original) != digest:
            raise CoursewareError('The file hash changed after copying.')
        pdf_path = original if kind == 'pdf' else stage / 'rendered.pdf'
        converter = 'direct_pdf'
        pptx_text = []
        if kind == 'pptx':
            pptx_text = _pptx_text(original)
        if kind != 'pdf':
            converter = _convert_powerpoint(original, pdf_path, stage)
        import pymupdf as fitz
        pages, warnings = [], []
        with fitz.open(pdf_path) as pdf:
            if pdf.needs_pass:
                raise CoursewareError('Encrypted courseware is not supported.')
            if not 0 < len(pdf) <= MAX_PAGES:
                raise CoursewareError('Courseware must contain 1–200 pages.')
            if pptx_text and len(pptx_text) != len(pdf):
                raise CoursewareError('The PPTX text page count does not match the rendered page count.')
            budget = min(MAX_PAGE_PIXELS, MAX_TOTAL_PIXELS // len(pdf))
            for number, page in enumerate(pdf, 1):
                width, height = page.rect.width, page.rect.height
                if not math.isfinite(width * height) or width <= 0 or height <= 0:
                    raise CoursewareError('Invalid courseware page dimensions.')
                scale = min(RENDER_LONG_EDGE / max(width, height), math.sqrt(budget / (width * height)))
                # A small guard keeps rounded raster dimensions below the pixel limit.
                scale *= 0.998
                picture = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                if picture.width * picture.height > budget:
                    raise CoursewareError('Courseware rendering exceeds the pixel limit.')
                name = f'page-{number:03d}.png'
                picture.save(stage / name)
                # Keep raw extraction for font-map diagnostics: geometric sorting
                # can discard control characters that reveal a broken encoding.
                native = (pptx_text[number - 1] if pptx_text else page.get_text('text', sort=False)).strip()
                text_source = 'pptx_native' if pptx_text else 'pdf_native'
                text = native
                native_quality = _native_text_quality(native)
                if native_quality != 'usable':
                    try:
                        recognized = _ocr(stage / name)
                        if recognized and (native_quality == 'broken_font_map' or
                                           len(''.join(recognized.split())) > len(''.join(native.split()))):
                            text, text_source = recognized, 'tesseract_chi_sim_eng'
                        elif native_quality == 'broken_font_map':
                            text, text_source = '', 'unreadable_font_map'
                            warnings.append(f'Page {number} has a damaged native font map and OCR returned no text. The page image was retained.')
                    except Exception as exc:
                        warnings.append(f'OCR did not complete for page {number} ({type(exc).__name__}). Native text was retained.')
                        if native_quality == 'broken_font_map':
                            text, text_source = '', 'unreadable_font_map'
                relative_image = (relative_root / name).as_posix()
                pages.append({'id': f'P{digest[:16]}_{number:03d}', 'document_id': document_id,
                              'page_number': number, 'text': text, 'text_source': text_source,
                              'native_text': native, 'native_text_quality': native_quality, 'image_path': relative_image,
                              'image_url': f'/api/review/media/{lesson_id}/{quote(relative_image)}',
                              'width': picture.width, 'height': picture.height})
        original_relative = (relative_root / original.name).as_posix()
        document = {'document_id': document_id, 'filename': source.name, 'sha256': digest,
                    'format': kind, 'source_url': f'/api/review/media/{lesson_id}/{quote(original_relative)}',
                    'source_path': original_relative, 'pages': pages, 'warnings': warnings,
                    'provenance': {'method': 'source_preserved_native_render', 'converter': converter,
                                   'importer_version': IMPORTER_VERSION,
                                   'renderer': 'PyMuPDF', 'renderer_version': fitz.VersionBind,
                                   'source_sha256': digest, 'rendered_pdf_sha256': _hash(pdf_path),
                                   'created_at': datetime.now(timezone.utc).isoformat(),
                                   'ocr': 'tesseract chi_sim+eng for image-only pages',
                                   'limits': {'bytes': MAX_DOCUMENT_BYTES, 'pages': MAX_PAGES,
                                              'page_pixels': MAX_PAGE_PIXELS, 'total_pixels': MAX_TOTAL_PIXELS},
                                   'notice': 'The courseware source file and page content were saved; text recognition and lesson matching still require review.'}}
        (stage / 'document.json').write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding='utf-8')
        stage.rename(root)
        return document
    finally:
        if stage.exists():
            resolved_stage, resolved_parent = stage.resolve(), parent.resolve()
            if resolved_stage == resolved_parent or not resolved_stage.is_relative_to(resolved_parent):
                raise CoursewareError('Refusing to clean a path outside the courseware working directory.')
            shutil.rmtree(stage)
