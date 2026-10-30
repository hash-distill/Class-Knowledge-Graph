"""Package one observed video page plus an untimed source note; never invent slides."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
SOURCE_SHA256='bd955b05ac65bcae85890ecda939e8e027389f3e739f1ea3775299db5231df1d'
sys.path.insert(0,str(ROOT))
from review.asr import inspect_video


def digest(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--video',type=Path,required=True)
    p.add_argument('--at-sec',type=float,default=45,choices=[45.0],help='This audited recipe packages only the 45-second projection of the fixed source')
    p.add_argument('--output',type=Path,default=ROOT/'artifacts/unified-demo-courseware')
    p.add_argument('--observed-text',type=Path,help='Same-source slides.json for a searchable OCR layer on the observed image')
    a=p.parse_args()
    import cv2
    import pymupdf as fitz
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    source=inspect_video(a.video)
    if source['sha256']!=SOURCE_SHA256:raise ValueError('This chemistry courseware recipe requires the audited original video SHA-256')
    searchable=''
    if a.observed_text:
        observed=json.loads(a.observed_text.read_text(encoding='utf-8'))
        if observed['provenance']['source_video']['sha256']!=source['sha256']:raise ValueError('Searchable OCR must come from this exact source video')
        matching=[row for row in observed['slides'] if row['start_sec']<=a.at_sec<row['end_sec'] and row.get('method')=='video_ocr']
        if len(matching)!=1:raise ValueError('The observed time must have one actual video OCR page')
        searchable=matching[0]['text']
    if not 0<=a.at_sec<source['duration_sec']:raise ValueError('Requested frame is outside source video')
    a.output.mkdir(parents=True,exist_ok=True)
    capture=cv2.VideoCapture(str(a.video));capture.set(cv2.CAP_PROP_POS_MSEC,a.at_sec*1000)
    ok,frame=capture.read();decoded_sec=capture.get(cv2.CAP_PROP_POS_MSEC)/1000;capture.release()
    if not ok or abs(decoded_sec-a.at_sec)>.1:raise ValueError('Source frame decode did not match requested time')
    image=a.output/'observed-page-45s.png'
    cv2.imencode('.png',frame)[1].tofile(str(image))
    title='Observed courseware provenance'
    rows=[
        ('Lesson','Chinese lesson: introduction to amino acids'),
        ('Frame source','Actual projected page in the original video, observed at 45 s'),
        ('Review','Compare speech, courseware, and visible behavior on the source timeline'),
        ('Provenance','Compiled from video frames; not the original author-supplied courseware'),
        ('Time limit','This note is absent from the video and receives no lesson timestamp'),
    ]
    pptx_path=a.output/'amino-acids-observed-courseware.pptx'
    deck=Presentation();deck.slide_width=Inches(12.8);deck.slide_height=Inches(7.2)
    deck.core_properties.title='Introduction to amino acids: observed video courseware'
    deck.core_properties.author='';deck.core_properties.subject='Same-source lesson demonstration; not original courseware'
    page=deck.slides.add_slide(deck.slide_layouts[6]);page.shapes.add_picture(str(image),0,0,deck.slide_width,deck.slide_height)
    page.notes_slide.notes_text_frame.text=f'Original video SHA-256: {source["sha256"]}; observed time: {a.at_sec} s; projected content is not cropped or altered.'
    note=deck.slides.add_slide(deck.slide_layouts[6])
    note.background.fill.solid();note.background.fill.fore_color.rgb=RGBColor.from_string('FCFCF7')
    def text_box(text,x,y,w,h,size,color='203C4A',bold=False):
        shape=note.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shape.text_frame;tf.word_wrap=True
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        run=tf.paragraphs[0].add_run();run.text=text;run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    text_box(title,.65,.55,11.5,.65,30,bold=True)
    text_box('Keep original frames so every evidence item can be replayed.',.65,1.35,11.5,.4,17,'698173')
    for i,(label,value) in enumerate(rows):
        y=2.15+i*.72;text_box(label,.7,y,1.4,.45,17,'698173');text_box(value,2.35,y,9.4,.48,19)
    text_box('Video hash and build record are saved. Page order is not lesson time.',.7,6.37,11.3,.35,13,'698173')
    deck.save(pptx_path)
    pdf_path=a.output/'amino-acids-observed-courseware.pdf'
    document=fitz.open();page=document.new_page(width=1280,height=720)
    page.insert_image(page.rect,filename=str(image))
    if searchable:
        # Searchable-scan text, verbatim from actual OCR. Original pixels stay intact.
        page.insert_text((60,105),searchable,fontname='china-s',fontsize=20,render_mode=3)
    page=document.new_page(width=1280,height=720);page.draw_rect(page.rect,fill=(.988,.988,.969),color=None)
    page.insert_text((65,95),title,fontname='china-s',fontsize=31,color=(.125,.235,.29))
    page.insert_text((65,145),'Keep original frames so every evidence item can be replayed.',fontname='china-s',fontsize=18,color=(.41,.51,.45))
    for i,(label,value) in enumerate(rows):
        y=244+i*72
        page.insert_text((70,y),label,fontname='china-s',fontsize=17,color=(.41,.51,.45))
        page.insert_text((235,y),value,fontname='china-s',fontsize=19,color=(.125,.235,.29))
    page.insert_text((70,658),'Video hash and build record are saved. Page order is not lesson time.',fontname='china-s',fontsize=13,color=(.41,.51,.45))
    document.set_metadata({'title':'Introduction to amino acids: observed video courseware','author':'','subject':'Compiled from observed source-video pages; not original courseware'})
    document.save(pdf_path)
    for i,page in enumerate(document):page.get_pixmap(matrix=fitz.Matrix(1,1)).save(a.output/f'preview-{i+1}.png')
    document.close()
    manifest={'source_video_sha256':source['sha256'],'source_duration_sec':source['duration_sec'],
              'observed_time_sec':a.at_sec,'decoded_time_sec':decoded_sec,
              'derived_from_source_video':True,'original_author_courseware':False,
              'page_roles':['observed_projection_frame','untimed_source_note'],
              'searchable_text_layer':{'source':'same_video_observed_ocr' if searchable else None,
                                       'text':searchable,'ocr_file_sha256':digest(a.observed_text) if searchable else None},
              'created_at':datetime.now(timezone.utc).isoformat(),
              'artifacts':{p.name:digest(p) for p in (image,pptx_path,pdf_path)}}
    (a.output/'provenance.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pptx':str(pptx_path.resolve()),'pdf':str(pdf_path.resolve()),'pages':2,'source_sha256':source['sha256']},ensure_ascii=False))


if __name__=='__main__':main()
